from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import tempfile
from collections.abc import Sequence
from pathlib import Path

from genbase_rag.embedding import create_text_embedding
from genbase_rag.models import KnowledgeEntry
from genbase_rag.text import extract_urls, normalize_text

DEFAULT_COLLECTION = "genbase_qa"
DEFAULT_EMBEDDING_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"


def create_entry_id(question: str, answer: str) -> str:
    normalized_pair = f"{normalize_text(question).casefold()}\0{normalize_text(answer)}"
    digest = hashlib.sha256(normalized_pair.encode("utf-8")).hexdigest()[:16].upper()
    return f"GBQA-{digest}"


def load_workbook_entries(source: Path, version: str = "") -> list[KnowledgeEntry]:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook.active
    header = [normalize_text(value).casefold() for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    try:
        question_index = header.index("question")
        answer_index = header.index("answer")
    except ValueError as exc:
        raise ValueError("The workbook must contain 'question' and 'answer' columns.") from exc

    entries: list[KnowledgeEntry] = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        question = normalize_text(row[question_index] if question_index < len(row) else "")
        answer = normalize_text(row[answer_index] if answer_index < len(row) else "")
        if not question and not answer:
            continue
        if not question or not answer:
            raise ValueError(f"Row {row_number} must contain both a question and an answer.")
        urls = extract_urls(answer)
        entries.append(
            KnowledgeEntry(
                id=create_entry_id(question, answer),
                question=question,
                answer=answer,
                source_url=urls[0] if urls else "",
                version=version,
            )
        )
    return entries


def write_jsonl(entries: Sequence[KnowledgeEntry], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="\n") as stream:
        for entry in entries:
            stream.write(json.dumps(entry.to_dict(), ensure_ascii=False) + "\n")


def rebuild_vector_index(entries: Sequence[KnowledgeEntry], destination: Path, model_name: str, collection_name: str = DEFAULT_COLLECTION) -> None:
    import chromadb

    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=destination.parent))
    backup = destination.with_name(f".{destination.name}.backup")
    try:
        model = create_text_embedding(model_name)
        documents = [entry.document_text for entry in entries]
        embeddings = [embedding.tolist() for embedding in model.embed(documents)]
        client = chromadb.PersistentClient(path=str(staging))
        collection = client.create_collection(collection_name, metadata={"hnsw:space": "cosine"})
        collection.add(
            ids=[entry.id for entry in entries],
            documents=documents,
            embeddings=embeddings,
            metadatas=[
                {
                    "question": entry.question,
                    "answer": entry.answer,
                    "category": entry.category,
                    "keywords": ",".join(entry.keywords),
                    "source_url": entry.source_url,
                    "language": entry.language,
                    "version": entry.version,
                }
                for entry in entries
            ],
        )
        if collection.count() != len(entries):
            raise RuntimeError("The Chroma index entry count does not match the source data.")
        client.close()

        if backup.exists():
            shutil.rmtree(backup)
        if destination.exists():
            destination.rename(backup)
        try:
            staging.rename(destination)
        except Exception:
            if backup.exists() and not destination.exists():
                backup.rename(destination)
            raise
        if backup.exists():
            shutil.rmtree(backup)
    finally:
        if staging.exists():
            shutil.rmtree(staging)


def parse_args(args: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the GenBase FAQ vector index.")
    parser.add_argument("--source", type=Path, required=True, help="Path to the GenBase QA workbook.")
    parser.add_argument("--output-dir", type=Path, required=True, help="Directory for normalized JSONL and Chroma data.")
    parser.add_argument("--version", default="", help="Knowledge-base version or source date.")
    parser.add_argument("--model", default=DEFAULT_EMBEDDING_MODEL, help="FastEmbed model name.")
    return parser.parse_args(args)


def main(args: Sequence[str] | None = None) -> None:
    options = parse_args(args)
    entries = load_workbook_entries(options.source, version=options.version)
    jsonl_path = options.output_dir / "genbase_qa.jsonl"
    staged_jsonl_path = options.output_dir / ".genbase_qa.jsonl.tmp"
    write_jsonl(entries, staged_jsonl_path)
    rebuild_vector_index(entries, options.output_dir / "chroma", options.model)
    staged_jsonl_path.replace(jsonl_path)
    print(f"Indexed {len(entries)} GenBase QA entries in {options.output_dir}")


if __name__ == "__main__":
    main()
